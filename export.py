# export.py
# -*- coding: utf-8 -*-
"""Экспорт отчётов в XLSX/PDF (A05: защита от formula injection)."""
import os
import datetime
from io import BytesIO

from flask import send_file
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

from validators import sanitize_cell
from reports import build_report, ALLOWED_KINDS


_REPORT_FONT = None


def _cyr_font():
    global _REPORT_FONT
    if _REPORT_FONT:
        return _REPORT_FONT
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
        "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/DejaVuSans.ttf",
    ]
    for path in candidates:
        if os.path.exists(path):
            try:
                pdfmetrics.registerFont(TTFont("CyrFont", path))
                _REPORT_FONT = "CyrFont"
                return _REPORT_FONT
            except Exception:
                continue
    _REPORT_FONT = "Helvetica"
    return _REPORT_FONT


def export_xlsx(title, headers, data, kind):
    wb = Workbook()
    ws = wb.active
    ws.title = (title or "Отчёт")[:31]
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append([sanitize_cell(h) for h in headers])
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=3, column=c)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in data:
        ws.append([sanitize_cell(c) for c in row])
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(str(h)) + 4)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"report_{kind}_{datetime.datetime.now():%Y%m%d_%H%M%S}.xlsx"
    return send_file(
        buf,
        as_attachment=True,
        download_name=fname,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def export_pdf(title, headers, data, kind):
    font = _cyr_font()
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24,
        title=title,
    )
    styles = getSampleStyleSheet()
    title_style = styles["Title"]
    title_style.fontName = font

    elements = [Paragraph(title, title_style), Spacer(1, 12)]
    table_data = [list(headers)] + [[str(c) for c in row] for row in data]
    t = Table(table_data, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f3a5f")),
        ("TEXTCOLOR",  (0, 0), (-1, 0), colors.white),
        ("FONTNAME",   (0, 0), (-1, -1), font),
        ("FONTSIZE",   (0, 0), (-1, -1), 9),
        ("GRID",       (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
        ("VALIGN",     (0, 0), (-1, -1), "MIDDLE"),
        ("PADDING",    (0, 0), (-1, -1), 4),
    ]))
    elements.append(t)
    doc.build(elements)
    buf.seek(0)
    fname = f"report_{kind}_{datetime.datetime.now():%Y%m%d_%H%M%S}.pdf"
    return send_file(buf, as_attachment=True, download_name=fname, mimetype="application/pdf")


def send_report(kind, fmt, date_from, date_to):
    if kind not in ALLOWED_KINDS:
        kind = "stock"
    if fmt not in ("pdf", "xlsx"):
        fmt = "pdf"
    title, headers, data = build_report(kind, date_from, date_to)
    if fmt == "xlsx":
        return export_xlsx(title, headers, data, kind)
    return export_pdf(title, headers, data, kind)